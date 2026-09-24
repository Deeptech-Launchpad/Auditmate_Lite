"""Recognising the sheets of a reporting pack, and the firm's own particulars.

Accounting systems export a "financial reporting" pack as one workbook: a
compilation report, a directory, an approval page, the statements, the notes.
Each sheet is already drawn FROM the books, so none of them is a source for
the accounts - but they read like one. A sheet with income and expense rows
and a debit or a credit was filed as a trial balance, and once the workbook
was split into files they all carried the same name, so the name could not
tell them apart either. What identifies a sheet is its title, on its first
row.

The compilation report is the one sheet worth reading: it is signed by the
practitioner, whose name and address SSRS 4410 requires and which no other
document gives. Most exports leave those as "[Insert field: accountants]",
and a placeholder is never taken as a name.
"""
import logging
import re
from pathlib import Path

log = logging.getLogger(__name__)

COMPILATION = "compilation_report"
PACK = "reporting_pack"

# The titles such a pack gives its sheets. Only these: a trial balance or a
# profit and loss exported on its own is titled differently and stays a source.
_PACK_TITLES = (
    "directory",
    "approval of financial report",
    "statement of comprehensive income",
    "statement of financial position",
    "statement of changes in equity",
    "statement of cash flows",
    "notes to the financial statements",
)

_PLACEHOLDER = re.compile(r"\[[^\]]*\]|\binsert\b|^\W*$", re.IGNORECASE)
_DATED = re.compile(r"^\s*dated\b", re.IGNORECASE)


def _cells(path, limit=60):
    """The non-empty text cells of a workbook's first sheet, in order."""
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        sheet = workbook.worksheets[0]
        out = []
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if value not in (None, ""):
                    out.append(" ".join(str(value).split()))
            if len(out) >= limit:
                break
        return out
    finally:
        workbook.close()


def sheet_kind(path):
    """COMPILATION, PACK, or None, from the title on the sheet's first row."""
    if Path(str(path)).suffix.lower() not in (".xlsx", ".xlsm"):
        return None
    try:
        cells = _cells(path, limit=1)
    except Exception:                                      # noqa: BLE001
        log.exception("Could not read the title of %s", path)
        return None
    if not cells:
        return None
    title = cells[0].lower()
    if title == "compilation report":
        return COMPILATION
    if any(title.startswith(t) for t in _PACK_TITLES):
        return PACK
    return None


def read_practitioner(path):
    """(name, address) written at the foot of a compilation report, or (None, None).

    The lines between the last paragraph and "Dated:". A line still holding a
    template placeholder - "[Insert field: accountants]" - is not a name.
    """
    try:
        cells = _cells(path, limit=200)
    except Exception:                                      # noqa: BLE001
        log.exception("Could not read the compilation report %s", path)
        return None, None

    tail = []
    for cell in reversed(cells):
        if _DATED.match(cell):
            tail = []
            continue
        if len(cell) > 90:                       # a paragraph of the report
            break
        tail.append(cell)
    tail.reverse()

    real = [c for c in tail if not _PLACEHOLDER.search(c)]
    if not real or len(real) != len(tail):
        # Any placeholder left means the firm never filled it in; taking the
        # lines that happen to be real would pair a name with a stranger's
        # address.
        return None, None
    return real[0], ", ".join(real[1:]) or None


def remember_practitioner(path, user_id=None):
    """Store what the report names as the firm's particulars, if unset.

    Never over a value somebody typed. Returns (name, address, message).
    """
    from . import disclosure_settings

    name, address = read_practitioner(path)
    if not name:
        return None, None, ("The practitioner's name and address are still "
                            "blank placeholders in this file, so nothing was "
                            "read from it")
    current = disclosure_settings.firm_defaults()
    values = {}
    if not current.get("practitioner_name"):
        values["practitioner_name"] = name
    if address and not current.get("practitioner_address"):
        values["practitioner_address"] = address
    if values:
        disclosure_settings.save(values, customer_id=None, user_id=user_id,
                                 commit=False)
    return name, address, "Practitioner details read from the compilation report"
